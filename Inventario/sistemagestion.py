import csv
import json
import os
from datetime import datetime
from typing import List, Dict
import Inventario.factura as IMPfactura
import Inventario.cliente as IMPcliente
import Inventario.productos as IMPproductos
import Inventario.venta as IMPventa
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_DISPONIBLE = True
except ImportError:
    EXCEL_DISPONIBLE = False
    print("Módulo openpyxl no disponible. Instale con: pip install openpyxl")
    
class SistemaGestionTienda:
    """Sistema principal de gestión de la tienda"""
    
    DESCUENTO_RESIDENT_EVIL = 0.15
    def __init__(self):
        self.productos: List[IMPproductos.Producto] = []
        self.clientes: Dict[str, IMPcliente.Cliente] = {}
        self.facturas: List[IMPfactura.Factura] = []
        self.contador_facturas = 1
        self.contador_clientes = 1
        self.cargar_productos_csv()
        self.cargar_datos_persistentes()


    def es_producto_resident_evil(self, producto: IMPproductos.Producto) -> bool:
        nombre_lower = producto.nombre.lower()
        keywords = ['resident evil', 'umbrella', 'stars', 'leon', 'jill', 'chris', 
                   'claire', 'ada wong', 'wesker', 'ethan', 'dimitrescu', 'nemesis']
        return any(keyword in nombre_lower for keyword in keywords)
    

    def obtener_precio_con_descuento(self, producto: IMPproductos.Producto) -> tuple:
        if self.es_producto_resident_evil(producto):
            descuento = producto.precio * self.DESCUENTO_RESIDENT_EVIL
            precio_final = producto.precio - descuento
            return (precio_final, descuento, True)
        return (producto.precio, 0.0, False)
    
    def cargar_datos_persistentes(self):
        """Carga clientes y facturas desde archivos JSON"""
        try:
            utils_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'utils'))
            clientes_path = os.path.join(utils_dir, 'clientes.json')
            # Cargar clientes
            if os.path.exists(clientes_path):
                with open(clientes_path, 'r', encoding='utf-8') as f:
                    datos_clientes = json.load(f)
                    for id_cliente, datos in datos_clientes.items():
                        cliente = IMPcliente.Cliente(
                            datos['id_cliente'],
                            datos['nombre'],
                            datos['email'],
                            datos['telefono']
                        )
                        cliente.historial_compras = datos.get('historial_compras', [])
                        self.clientes[id_cliente] = cliente
                    
                    # Actualizar contador
                    if self.clientes:
                        ultimo_id = max([int(c.id_cliente[3:]) for c in self.clientes.values()])
                        self.contador_clientes = ultimo_id + 1
                
                print(f"{len(self.clientes)} clientes cargados")
                
            facturas_path = os.path.join(utils_dir, 'facturas.json')
            # Cargar facturas
            if os.path.exists(facturas_path):
                with open(facturas_path, 'r', encoding='utf-8') as f:
                    datos_facturas = json.load(f)
                    for datos_factura in datos_facturas:
                        cliente = self.clientes.get(datos_factura['id_cliente'])
                        if cliente:
                            factura = IMPfactura.Factura(datos_factura['id_factura'], cliente)
                            factura.fecha = datetime.fromisoformat(datos_factura['fecha'])
                            factura.total = datos_factura['total']
                            
                            # Reconstruir items
                            for item_data in datos_factura['items']:
                                producto = self.buscar_producto(item_data['id_producto'])
                                if producto:
                                    item = IMPfactura.ItemFactura(producto, item_data['cantidad'])
                                    factura.items.append(item)
                            
                            self.facturas.append(factura)
                    
                    # Actualizar contador
                    if self.facturas:
                        ultimo_id = max([int(f.id_factura[1:]) for f in self.facturas])
                        self.contador_facturas = ultimo_id + 1
                
                print(f"{len(self.facturas)} facturas cargadas")
        
        except Exception as e:
            print(f"Error al cargar datos: {e}")
    
    def guardar_datos_persistentes(self):
        """Guarda clientes y facturas en archivos JSON"""
        try:
            utils_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'utils'))
            os.makedirs(utils_dir, exist_ok=True)
            # Guardar clientes
            datos_clientes = {}
            for id_cliente, cliente in self.clientes.items():
                datos_clientes[id_cliente] = {
                    'id_cliente': cliente.id_cliente,
                    'nombre': cliente.nombre,
                    'email': cliente.email,
                    'telefono': cliente.telefono,
                    'historial_compras': cliente.historial_compras
                }
            clientes_path = os.path.join(utils_dir, 'clientes.json')
            with open(clientes_path, 'w', encoding='utf-8') as f:
                json.dump(datos_clientes, f, indent=2, ensure_ascii=False)
            
            # Guardar facturas
            datos_facturas = []
            for factura in self.facturas:
                items_data = []
                for item in factura.items:
                    items_data.append({
                        'id_producto': item.producto.id_producto,
                        'cantidad': item.cantidad
                    })
                
                datos_facturas.append({
                    'id_factura': factura.id_factura,
                    'id_cliente': factura.cliente.id_cliente,
                    'fecha': factura.fecha.isoformat(),
                    'total': factura.total,
                    'items': items_data
                })
            facturas_path = os.path.join(utils_dir, 'facturas.json')
            with open(facturas_path, 'w', encoding='utf-8') as f:
                json.dump(datos_facturas, f, indent=2, ensure_ascii=False)
            
            print("Datos guardados exitosamente")
        
        except Exception as e:
            print(f"Error al guardar datos: {e}")
    
    def cargar_productos_csv(self):
        """Carga productos desde archivos CSV"""
        try:
            utils_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'utils'))
            videojuegos_path = os.path.join(utils_dir, 'videojuegos.csv')
            productos_gaming_path = os.path.join(utils_dir, 'productos_gaming.csv')
            # Cargar videojuegos
            with open(videojuegos_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    vj = IMPproductos.Videojuego(
                        row['id'], row['nombre'], float(row['precio']),
                        int(row['stock']), row['plataforma'], 
                        row['genero'], int(row['año'])
                    )
                    self.productos.append(vj)
            
            # Cargar productos gaming
            with open(productos_gaming_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                for row in reader:
                    pg = IMPproductos.ProductoGaming(
                        row['id'], row['nombre'], float(row['precio']),
                        int(row['stock']), row['tipo'], row['marca']
                    )
                    self.productos.append(pg)
            
            print("Productos cargados exitosamente")

            productos_re = sum(1 for p in self.productos if self.es_producto_resident_evil(p))
            if productos_re > 0:
                print(f" {productos_re} productos con descuento Resident Evil ({self.DESCUENTO_RESIDENT_EVIL*100:.0f}%)")
        except FileNotFoundError:
            print("Archivos CSV no encontrados. Creando archivos de ejemplo...")
            self.crear_csv_ejemplo()
    
  
    def mostrar_catalogo(self):
        """Muestra el catálogo completo de productos"""
        print("\n" + "="*80)
        print("CATÁLOGO DE PRODUCTOS".center(80))
        print("="*80)
        
        print("\n--- VIDEOJUEGOS ---")
        for p in self.productos:
            if isinstance(p, IMPproductos.Videojuego):
                print(p)
        
        print("\n--- PRODUCTOS GAMING ---")
        for p in self.productos:
            if isinstance(p, IMPproductos.ProductoGaming):
                print(p)
        
        print("="*80)
    
    def buscar_producto(self, id_producto: str) -> IMPproductos.Producto:
        """Busca un producto por ID"""
        for p in self.productos:
            if p.id_producto == id_producto:
                return p
        return None
    
    def buscar_producto_por_nombre(self, nombre: str) -> List[IMPproductos.Producto]:
        """Busca productos por nombre (coincidencia parcial)"""
        nombre_lower = nombre.lower()
        resultados = []
        for p in self.productos:
            if nombre_lower in p.nombre.lower():
                resultados.append(p)
        return resultados
    
    def mostrar_catalogo_compacto(self):
        """Muestra el catálogo en formato compacto para consulta rápida"""
        print("\n" + "="*100)
        print("CATÁLOGO RÁPIDO".center(100))
        print("="*100)
        
        print("\n--- VIDEOJUEGOS ---")
        print(f"{'ID':<8} {'Nombre':<35} {'Precio':<12} {'Plataforma':<12} {'Stock':<8}")
        print("-"*100)
        for p in self.productos:
            if isinstance(p, IMPproductos.Videojuego):
                precio_final, descuento, tiene_desc = self.obtener_precio_con_descuento(p)
                if tiene_desc:
                    precio_str = f"${precio_final:.2f} "
                else:
                    precio_str = f"${p.precio:.2f}"
                print(f"{p.id_producto:<8} {p.nombre:<40} {precio_str:<15} {p.plataforma:<12} {p.stock:<8}")
        
        print("\n--- PRODUCTOS GAMING ---")
        print(f"{'ID':<8} {'Nombre':<35} {'Precio':<12} {'Tipo':<12} {'Stock':<8}")
        print("-"*100)
        for p in self.productos:
            if isinstance(p, IMPproductos.ProductoGaming):
                print(f"{p.id_producto:<8} {p.nombre:<35} ${p.precio:<11.2f} {p.tipo:<12} {p.stock:<8}")
        print("="*100)
    
    def registrar_cliente(self):
        """Registra un nuevo cliente"""
        print("\n--- REGISTRO DE CLIENTE ---")
        while True:
                nombre = input("Nombre completo: ").strip()
                if nombre:
                    break
                print("El nombre no puede estar vacío. Intente de nuevo.")

        while True:
            email = input("Email: ").strip()
            if "@" in email and "." in email:
                break
            print("Email inválido. Intente de nuevo.")

        while True:
            telefono = input("Teléfono: ").strip()
            if telefono.isdigit() and len(telefono) >= 7:
                break
            print("El teléfono debe contener solo números y tener al menos 7 dígitos.")
        id_cliente = f"CLI{self.contador_clientes:04d}"
        cliente = IMPcliente.Cliente(id_cliente, nombre, email, telefono)
        self.clientes[id_cliente] = cliente
        self.contador_clientes += 1
        
        self.guardar_datos_persistentes()
        
        print(f"Cliente registrado exitosamente. ID: {id_cliente}")
        return cliente
    
    def buscar_cliente_por_nombre(self, nombre: str) -> List[IMPcliente.Cliente]:
        """Busca clientes por nombre (coincidencia parcial)"""
        nombre_lower = nombre.lower()
        resultados = []
        for cliente in self.clientes.values():
            if nombre_lower in cliente.nombre.lower():
                resultados.append(cliente)
        return resultados
    
    def mostrar_clientes_compacto(self):
        """Muestra lista de clientes en formato compacto"""
        print("\n" + "="*100)
        print("LISTA DE CLIENTES".center(100))
        print("="*100)
        print(f"{'ID':<10} {'Nombre':<30} {'Email':<30} {'Teléfono':<15}")
        print("-"*100)
        for cliente in self.clientes.values():
            print(f"{cliente.id_cliente:<10} {cliente.nombre:<30} {cliente.email:<30} {cliente.telefono:<15}")
        print("="*100)
    
    def seleccionar_cliente(self) -> IMPcliente.Cliente:
        """Interfaz mejorada para seleccionar o registrar cliente"""
        print("\n--- SELECCIÓN DE CLIENTE ---")
        print("Opciones: [ID] / [buscar:nombre] / [lista] / [nuevo]")
        entrada = input("Ingrese opción: ").strip()
        
        if entrada.lower() == 'nuevo':
            return self.registrar_cliente()
        
        elif entrada.lower() == 'lista':
            self.mostrar_clientes_compacto()
            entrada = input("\nAhora ingrese ID o buscar:nombre: ").strip()
        
        if entrada.lower().startswith('buscar:'):
            # Búsqueda por nombre
            nombre_busqueda = entrada[7:].strip()
            resultados = self.buscar_cliente_por_nombre(nombre_busqueda)
            
            if not resultados:
                print("No se encontraron clientes con ese nombre")
                print("¿Desea registrar un nuevo cliente? (s/n)")
                if input().lower() == 's':
                    return self.registrar_cliente()
                return None
            
            print(f"\n--- Resultados de búsqueda: '{nombre_busqueda}' ---")
            for i, c in enumerate(resultados, 1):
                print(f"{i}. {c}")
            
            try:
                seleccion = int(input("\nSeleccione número (0 para cancelar): "))
                if seleccion == 0:
                    return None
                return resultados[seleccion - 1]
            except (ValueError, IndexError):
                print("✗ Selección inválida")
                return None
        else:
            # Búsqueda por ID
            cliente = self.clientes.get(entrada)
            if not cliente:
                print("Cliente no encontrado")
                print("Tip: Use 'lista' para ver todos los clientes o 'buscar:nombre' para buscar")
                print("¿Desea registrar un nuevo cliente? (s/n)")
                if input().lower() == 's':
                    return self.registrar_cliente()
                return None
            return cliente
    
    def crear_factura(self):
        """Crea una nueva factura"""
        print("\n--- CREAR FACTURA ---")
        
        # Seleccionar o registrar cliente con interfaz mejorada
        cliente = self.seleccionar_cliente()
        if not cliente:
            print(" Operación cancelada")
            return
        
        # Crear factura
        id_factura = f"F{self.contador_facturas:05d}"
        factura = IMPfactura.Factura(id_factura, cliente)
        self.contador_facturas += 1
        
        # Agregar productos
        while True:
            print("\n--- Agregar producto ---")
            print("Opciones: [ID] / [buscar:nombre] / [catalogo] / [fin]")
            entrada = input("Ingrese opción: ").strip()
            
            if entrada.lower() == 'fin':
                break
            
            elif entrada.lower() == 'catalogo':
                self.mostrar_catalogo_compacto()
                continue
            
            elif entrada.lower().startswith('buscar:'):
                # Búsqueda por nombre
                nombre_busqueda = entrada[7:].strip()
                resultados = self.buscar_producto_por_nombre(nombre_busqueda)
                
                if not resultados:
                    print("No se encontraron productos con ese nombre")
                    continue
                
                print(f"\n--- Resultados de búsqueda: '{nombre_busqueda}' ---")
                for i, p in enumerate(resultados, 1):
                    print(f"{i}. {p}")
                
                try:
                    seleccion = int(input("\nSeleccione número (0 para cancelar): "))
                    if seleccion == 0:
                        continue
                    producto = resultados[seleccion - 1]
                except (ValueError, IndexError):
                    print("Selección inválida")
                    continue
            else:
                # Búsqueda por ID
                producto = self.buscar_producto(entrada)
                if not producto:
                    print("Producto no encontrado")
                    print("Tip: Use 'catalogo' para ver todos los productos o 'buscar:nombre' para buscar")
                    continue
            
        
        precio_final, descuento, tiene_desc = self.obtener_precio_con_descuento(producto)
        print(f"\nProducto seleccionado: {producto.nombre}")
        if tiene_desc:
            print(f"Precio original: ${producto.precio:.2f}")
            print(f"Precio con descuento RE (15%): ${precio_final:.2f}")
            print(f"Ahorro: ${descuento:.2f}")
        else:
            print(f"Precio: ${producto.precio:.2f}")
        print(f"Stock disponible: {producto.stock}")
        
        try:
            cantidad = int(input("Cantidad: "))
            
            # Aplicar descuento si corresponde
            if tiene_desc:
                producto_temp = type(producto).__new__(type(producto))
                producto_temp.__dict__.update(producto.__dict__)
                producto_temp.precio = precio_final
                
                if factura.agregar_item(producto_temp, cantidad):
                    print(f" Producto agregado con descuento RE (Ahorro: ${descuento * cantidad:.2f})")
                else:
                    print("Stock insuficiente")
            else:
                if factura.agregar_item(producto, cantidad):
                    print(" Producto agregado a la factura")
                else:
                    print("Stock insuficiente")
        except ValueError:
            print("Cantidad inválida")


        if factura.items:
            self.facturas.append(factura)
            # Guardar automáticamente
            self.guardar_datos_persistentes()
            print(factura.generar_resumen())
            print("Factura creada exitosamente")
        else:
            print("No se agregaron productos. Factura cancelada.")
    
    def exportar_factura(self, id_factura: str = None):
        """Exporta una factura a archivo de texto"""
        if not id_factura:
            # Mostrar lista de facturas
            if not self.facturas:
                print("No hay facturas para exportar")
                return
            
            print("\n--- LISTA DE FACTURAS ---")
            print(f"{'#':<4} {'ID':<12} {'Cliente':<30} {'Fecha':<20} {'Total':<12}")
            print("-"*85)
            for i, factura in enumerate(self.facturas, 1):
                print(f"{i:<4} {factura.id_factura:<12} {factura.cliente.nombre:<30} "
                      f"{factura.fecha.strftime('%Y-%m-%d %H:%M'):<20} ${factura.total:<11.2f}")
            
            seleccion = input("\nIngrese número o ID de factura (0 para cancelar): ").strip()
            if seleccion == '0':
                return
            
            # Intentar como número primero
            try:
                idx = int(seleccion) - 1
                if 0 <= idx < len(self.facturas):
                    id_factura = self.facturas[idx].id_factura
            except ValueError:
                id_factura = seleccion
        
        for factura in self.facturas:
            if factura.id_factura == id_factura:
                export_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'facturas'))
                os.makedirs(export_dir, exist_ok=True)
                nombre_archivo = f"factura_{id_factura}.txt"
                ruta_completa = os.path.join(export_dir, nombre_archivo)
                with open(ruta_completa, 'w', encoding='utf-8') as f:
                    f.write(factura.generar_resumen())
                print(f"✓ Factura exportada a {ruta_completa}")
                return
        print("Factura no encontrada")
    
    def exportar_factura_excel(self, id_factura: str = None):
        """Exporta una factura individual a Excel"""
        if not EXCEL_DISPONIBLE:
            print(" Módulo openpyxl no disponible. Instale con: pip install openpyxl")
            return
        
        if not id_factura:
            # Mostrar lista de facturas
            if not self.facturas:
                print(" No hay facturas para exportar")
                return
            
            print("\n--- LISTA DE FACTURAS ---")
            print(f"{'#':<4} {'ID':<12} {'Cliente':<30} {'Fecha':<20} {'Total':<12}")
            print("-"*85)
            for i, factura in enumerate(self.facturas, 1):
                print(f"{i:<4} {factura.id_factura:<12} {factura.cliente.nombre:<30} "
                      f"{factura.fecha.strftime('%Y-%m-%d %H:%M'):<20} ${factura.total:<11.2f}")
            
            seleccion = input("\nIngrese número o ID de factura (0 para cancelar): ").strip()
            if seleccion == '0':
                return
            
            # Intentar como número primero
            try:
                idx = int(seleccion) - 1
                if 0 <= idx < len(self.facturas):
                    id_factura = self.facturas[idx].id_factura
            except ValueError:
                id_factura = seleccion
        
        for factura in self.facturas:
            if factura.id_factura == id_factura:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = f"Factura {id_factura}"
                
                # Estilos
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True, size=12)
                border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                # Encabezado de la factura
                ws['A1'] = "FACTURA"
                ws['A1'].font = Font(bold=True, size=16)
                ws.merge_cells('A1:E1')
                ws['A1'].alignment = Alignment(horizontal='center')
                
                ws['A3'] = "Número de Factura:"
                ws['B3'] = factura.id_factura
                ws['A4'] = "Fecha:"
                ws['B4'] = factura.fecha.strftime('%Y-%m-%d %H:%M:%S')
                ws['A5'] = "Cliente:"
                ws['B5'] = factura.cliente.nombre
                ws['A6'] = "ID Cliente:"
                ws['B6'] = factura.cliente.id_cliente
                ws['A7'] = "Email:"
                ws['B7'] = factura.cliente.email
                ws['A8'] = "Teléfono:"
                ws['B8'] = factura.cliente.telefono
                
                # Estilo para encabezados de información
                for row in range(3, 9):
                    ws[f'A{row}'].font = Font(bold=True)
                
                # Tabla de productos
                row_inicio = 10
                headers = ['Producto', 'Cantidad', 'Precio Unitario', 'Subtotal']
                for col, header in enumerate(headers, 1):
                    cell = ws.cell(row=row_inicio, column=col, value=header)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.border = border
                    cell.alignment = Alignment(horizontal='center')
                
                # Datos de productos
                row_actual = row_inicio + 1
                for item in factura.items:
                    ws.cell(row=row_actual, column=1, value=item.producto.nombre).border = border
                    ws.cell(row=row_actual, column=2, value=item.cantidad).border = border
                    ws.cell(row=row_actual, column=3, value=item.producto.precio).border = border
                    ws.cell(row=row_actual, column=4, value=item.subtotal).border = border
                    
                    # Formato de moneda
                    ws.cell(row=row_actual, column=3).number_format = '$#,##0.00'
                    ws.cell(row=row_actual, column=4).number_format = '$#,##0.00'
                    
                    row_actual += 1
                
                # Total
                row_total = row_actual + 1
                ws.cell(row=row_total, column=3, value="TOTAL:").font = Font(bold=True, size=12)
                ws.cell(row=row_total, column=3).alignment = Alignment(horizontal='right')
                ws.cell(row=row_total, column=4, value=factura.total).font = Font(bold=True, size=12)
                ws.cell(row=row_total, column=4).number_format = '$#,##0.00'
                ws.cell(row=row_total, column=4).fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                
                # Ajustar ancho de columnas
                ws.column_dimensions['A'].width = 40
                ws.column_dimensions['B'].width = 15
                ws.column_dimensions['C'].width = 20
                ws.column_dimensions['D'].width = 15
                
                # Guardar archivo
                export_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'facturas'))
                os.makedirs(export_dir, exist_ok=True)
                nombre_archivo = f"factura_{id_factura}.xlsx"
                ruta_completa = os.path.join(export_dir, nombre_archivo)
                wb.save(ruta_completa)
                print(f"✓ Factura exportada a Excel: {ruta_completa}")
                return
        
        print("Factura no encontrada")
    
    def exportar_todas_facturas_excel(self):
        """Exporta todas las facturas a un archivo Excel con múltiples hojas"""
        if not EXCEL_DISPONIBLE:
            print(" Módulo openpyxl no disponible. Instale con: pip install openpyxl")
            return
        
        if not self.facturas:
            print(" No hay facturas para exportar")
            return
        
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remover hoja por defecto
        
        # Estilos comunes
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Crear hoja resumen
        ws_resumen = wb.create_sheet("Resumen")
        ws_resumen['A1'] = "RESUMEN DE FACTURAS"
        ws_resumen['A1'].font = Font(bold=True, size=14)
        ws_resumen.merge_cells('A1:F1')
        ws_resumen['A1'].alignment = Alignment(horizontal='center')
        
        headers_resumen = ['Factura', 'Fecha', 'Cliente', 'ID Cliente', 'Total', 'Items']
        for col, header in enumerate(headers_resumen, 1):
            cell = ws_resumen.cell(row=3, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center')
        
        # Llenar resumen
        row = 4
        total_general = 0
        for factura in self.facturas:
            ws_resumen.cell(row=row, column=1, value=factura.id_factura).border = border
            ws_resumen.cell(row=row, column=2, value=factura.fecha.strftime('%Y-%m-%d')).border = border
            ws_resumen.cell(row=row, column=3, value=factura.cliente.nombre).border = border
            ws_resumen.cell(row=row, column=4, value=factura.cliente.id_cliente).border = border
            ws_resumen.cell(row=row, column=5, value=factura.total).border = border
            ws_resumen.cell(row=row, column=6, value=len(factura.items)).border = border
            
            ws_resumen.cell(row=row, column=5).number_format = '$#,##0.00'
            total_general += factura.total
            row += 1
        
        # Total general
        ws_resumen.cell(row=row+1, column=4, value="TOTAL GENERAL:").font = Font(bold=True)
        ws_resumen.cell(row=row+1, column=5, value=total_general).font = Font(bold=True)
        ws_resumen.cell(row=row+1, column=5).number_format = '$#,##0.00'
        ws_resumen.cell(row=row+1, column=5).fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        
        # Ajustar columnas del resumen
        ws_resumen.column_dimensions['A'].width = 12
        ws_resumen.column_dimensions['B'].width = 12
        ws_resumen.column_dimensions['C'].width = 30
        ws_resumen.column_dimensions['D'].width = 12
        ws_resumen.column_dimensions['E'].width = 12
        ws_resumen.column_dimensions['F'].width = 8
        
        # Crear hojas individuales para cada factura
        for factura in self.facturas:
            ws = wb.create_sheet(f"{factura.id_factura}")
            
            # Encabezado
            ws['A1'] = "FACTURA"
            ws['A1'].font = Font(bold=True, size=14)
            ws.merge_cells('A1:D1')
            ws['A1'].alignment = Alignment(horizontal='center')
            
            ws['A3'] = "Número:"
            ws['B3'] = factura.id_factura
            ws['A4'] = "Fecha:"
            ws['B4'] = factura.fecha.strftime('%Y-%m-%d %H:%M')
            ws['A5'] = "Cliente:"
            ws['B5'] = factura.cliente.nombre
            ws['A6'] = "ID Cliente:"
            ws['B6'] = factura.cliente.id_cliente
            
            for row in range(3, 7):
                ws[f'A{row}'].font = Font(bold=True)
            
            # Tabla de productos
            row_inicio = 8
            headers_prod = ['Producto', 'Cantidad', 'Precio Unit.', 'Subtotal']
            for col, header in enumerate(headers_prod, 1):
                cell = ws.cell(row=row_inicio, column=col, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
            
            row_actual = row_inicio + 1
            for item in factura.items:
                ws.cell(row=row_actual, column=1, value=item.producto.nombre).border = border
                ws.cell(row=row_actual, column=2, value=item.cantidad).border = border
                ws.cell(row=row_actual, column=3, value=item.producto.precio).border = border
                ws.cell(row=row_actual, column=4, value=item.subtotal).border = border
                
                ws.cell(row=row_actual, column=3).number_format = '$#,##0.00'
                ws.cell(row=row_actual, column=4).number_format = '$#,##0.00'
                row_actual += 1
            
            # Total
            ws.cell(row=row_actual+1, column=3, value="TOTAL:").font = Font(bold=True)
            ws.cell(row=row_actual+1, column=4, value=factura.total).font = Font(bold=True)
            ws.cell(row=row_actual+1, column=4).number_format = '$#,##0.00'
            
            ws.column_dimensions['A'].width = 35
            ws.column_dimensions['B'].width = 12
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 15
        
        # Guardar archivo
        export_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'facturas'))
        os.makedirs(export_dir, exist_ok=True)
        nombre_archivo = f"facturas_completas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        ruta_completa = os.path.join(export_dir, nombre_archivo)
        wb.save(ruta_completa)
        print(f"Todas las facturas exportadas a Excel: {ruta_completa}")
        print(f"  - Total de facturas: {len(self.facturas)}")
        print(f"  - Total general: ${total_general:.2f}")
    
    def agregar_producto_inventario(self):
        """Agrega un nuevo producto al inventario"""
        print("\n--- AGREGAR PRODUCTO AL INVENTARIO ---")
        print("1. Videojuego")
        print("2. Producto Gaming")
        tipo = input("Tipo de producto: ")
        
        if tipo == "1":
            id_prod = input("ID: ")
            nombre = input("Nombre: ")
            precio = float(input("Precio: "))
            stock = int(input("Stock: "))
            plataforma = input("Plataforma: ")
            genero = input("Género: ")
            año = int(input("Año: "))
            
            producto = IMPproductos.Videojuego(id_prod, nombre, precio, stock, plataforma, genero, año)
        else:
            id_prod = input("ID: ")
            nombre = input("Nombre: ")
            precio = float(input("Precio: "))
            stock = int(input("Stock: "))
            tipo_prod = input("Tipo (Control/Auriculares/etc): ")
            marca = input("Marca: ")
            
            producto = IMPproductos.ProductoGaming(id_prod, nombre, precio, stock, tipo_prod, marca)
        
        self.productos.append(producto)
        print("Producto agregado exitosamente")
    
    def evaluar_compra_usado(self):
        from Inventario.venta import EvaluadorProducto, TipoProducto

        print("\n" + "="*80)
        print("EVALUACIÓN DE PRODUCTO USADO".center(80))
        print("="*80)
    
        print("\n¿Cómo desea iniciar la evaluación?")
        print("  1. Buscar producto en nuestro catálogo")
        print("  2. Evaluación manual (producto no está en catálogo)")
    
        opcion = input("\nSeleccione opción (1/2): ").strip()
    
        evaluador = EvaluadorProducto()
    

        if opcion != "1":
            evaluador.iniciar_evaluacion()
            return
    

        print("\n--- BUSCAR PRODUCTO ---")
        print("Opciones: [ID] / [buscar:nombre] / [catalogo]")
        entrada = input("Ingrese opción: ").strip()
    
        if entrada.lower() == 'catalogo':
            self.mostrar_catalogo_compacto()
            entrada = input("\nAhora ingrese ID o buscar:nombre: ").strip()
    
        producto_base = None
    
        # Búsqueda por nombre
        if entrada.lower().startswith('buscar:'):
            nombre_busqueda = entrada[7:].strip()
            resultados = self.buscar_producto_por_nombre(nombre_busqueda)
            
            if not resultados:
                print("✗ No se encontraron productos con ese nombre")
                print("\n¿Desea hacer evaluación manual? (s/n): ")
                if input().lower() == 's':
                    evaluador.iniciar_evaluacion()
                return
            
            print(f"\n--- Resultados de búsqueda: '{nombre_busqueda}' ---")
            for i, p in enumerate(resultados, 1):
                print(f"{i}. {p}")
            
            try:
                seleccion = int(input("\nSeleccione número (0 para cancelar): "))
                if seleccion == 0:
                    return
                producto_base = resultados[seleccion - 1]
            except (ValueError, IndexError):
                print("✗ Selección inválida")
                return
    
        # Búsqueda por ID
        else:
            producto_base = self.buscar_producto(entrada)
            
            if not producto_base:
                print("✗ Producto no encontrado")
                print("Tip: Use 'catalogo' para ver todos los productos")
                print("\n¿Desea hacer evaluación manual? (s/n): ")
                if input().lower() == 's':
                    evaluador.iniciar_evaluacion()
                return
    
        # ===== PRODUCTO ENCONTRADO =====
        if producto_base:
            print("\n" + "="*80)
            print("PRODUCTO DE REFERENCIA".center(80))
            print("="*80)
            print(f"\n✓ {producto_base.nombre}")
            print(f"  Precio nuevo en tienda: ${producto_base.precio:.2f}")
            print(f"  Categoría: {producto_base.categoria}")
            
            # Establecer información base
            evaluador.producto_nombre = producto_base.nombre
            evaluador.precio_nuevo = producto_base.precio
            
            # Seleccionar tipo de producto
            print("\n--- TIPO DE PRODUCTO ---")
            print("Seleccione el tipo para calcular depreciación correcta:")
            tipos = list(TipoProducto)
            for i, tipo in enumerate(tipos, 1):
                print(f"  {i}. {tipo.value.title()}")
            
            try:
                tipo_idx = int(input("\nSeleccione el tipo (número): ")) - 1
                evaluador.tipo_producto = tipos[tipo_idx]
                print(f"✓ Tipo seleccionado: {evaluador.tipo_producto.value.title()}")
            except:
                print("⚠ Selección inválida, usando 'Otros'")
    
    def menu_principal(self):
        """Menú principal del sistema"""
        while True:
            print("\n" + "="*80)
            print("SISTEMA DE GESTIÓN - TIENDA DE VIDEOJUEGOS".center(80))
            print("="*80)
            print("1. Ver catálogo de productos")
            print("2. Registrar cliente")
            print("3. Crear factura")
            print("4. Exportar factura (TXT)")
            print("5. Exportar factura (Excel)")
            print("6. Exportar TODAS las facturas (Excel)")
            print("7. Agregar producto al inventario")
            print("8. Evaluar compra de producto usado")
            print("9. Ver clientes registrados")
            print("10. Ver facturas")
            print("11. Salir")
            print("="*80)
            
            opcion = input("\nSeleccione una opción: ")
            
            if opcion == "1":
                self.mostrar_catalogo()
            elif opcion == "2":
                self.registrar_cliente()
            elif opcion == "3":
                self.crear_factura()
            elif opcion == "4":
                self.exportar_factura()
            elif opcion == "5":
                self.exportar_factura_excel()
            elif opcion == "6":
                self.exportar_todas_facturas_excel()
            elif opcion == "7":
                self.agregar_producto_inventario()
            elif opcion == "8":
                self.evaluar_compra_usado()
            elif opcion == "9":
                print("\n--- CLIENTES REGISTRADOS ---")
                if self.clientes:
                    self.mostrar_clientes_compacto()
                else:
                    print("No hay clientes registrados")
            elif opcion == "10":
                print("\n--- FACTURAS ---")
                if self.facturas:
                    print(f"{'ID':<12} {'Cliente':<30} {'Fecha':<20} {'Total':<12}")
                    print("-"*80)
                    for factura in self.facturas:
                        print(f"{factura.id_factura:<12} {factura.cliente.nombre:<30} "
                              f"{factura.fecha.strftime('%Y-%m-%d %H:%M'):<20} ${factura.total:<11.2f}")
                else:
                    print("No hay facturas registradas")
            elif opcion == "11":
                print("\n¿Desea guardar los cambios antes de salir? (s/n)")
                if input().lower() == 's':
                    self.guardar_datos_persistentes()
                print("\n¡Hasta pronto!")
                break
            else:
                print("Opción inválida")
            
            input("\nPresione Enter para continuar...")
